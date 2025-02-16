from docx import Document
from docx.shared import Pt  # برای تنظیم اندازه فونت
from docx.oxml.ns import qn  # برای پشتیبانی از زبان فارسی
from docx.oxml import OxmlElement

# Load the template Word file
template_path = 'template.docx'  # مسیر فایل نمونه شما
output_path = 'output.docx'  # مسیر ذخیره فایل خروجی

# داده‌های مثال
data = {"نام و نام خانوادگی": "سجاد محمدی", "نام پدر": "حسن", "کد ملی": "1245", "تاریخ تولد": "11/12/789", "ماه": "دی", "سال": "1400", "شهرستان": "تهران", "استان": "یزد", "پایه خدمتی": "1/10/403", "شماره": "4258", "تاریخ": "1/6/1403", "نام دوره": "مهارت چهارگانه"}

# فونت پیش‌فرض فارسی
default_font = "B Nazanin"

# Load the template file only once
doc = Document(template_path)

# Replace placeholders with actual values
for paragraph in doc.paragraphs:
    for key, value in data.items():  # استفاده از داده‌ها برای جایگزینی
        if f"{{{{{key}}}}}" in paragraph.text:
            # Update the text
            paragraph.text = paragraph.text.replace(f"{{{{{key}}}}}", value)
            
            # تنظیم فونت و استایل
            for run in paragraph.runs:
                run.font.name = default_font  # تنظیم فونت
                run.font.size = Pt(14)  # تنظیم اندازه فونت
                
                # پشتیبانی از زبان فارسی
                r = run._element
                rPr = r.get_or_add_rPr()
                rFonts = OxmlElement('w:rFonts')
                rFonts.set(qn('w:ascii'), default_font)
                rFonts.set(qn('w:hAnsi'), default_font)
                rFonts.set(qn('w:cs'), default_font)
                rPr.append(rFonts)

# Save the customized file
doc.save(output_path)
